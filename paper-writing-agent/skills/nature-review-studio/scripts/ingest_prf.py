#!/usr/bin/env python3
"""Extract explicitly supplied PRF material for a new task-local lesson draft.

No case graph, model training, automatic anonymization or global memory update.
"""
import argparse, json
from datetime import datetime, timezone
from pathlib import Path

def extract(path):
    suffix=path.suffix.lower()
    if suffix=='.pdf':
        from pypdf import PdfReader
        return [{'page':i,'text':p.extract_text() or ''} for i,p in enumerate(PdfReader(path).pages,1)]
    if suffix=='.docx':
        from docx import Document
        d=Document(path)
        return [{'block':i,'text':p.text} for i,p in enumerate(d.paragraphs,1) if p.text]+[{'table':i,'rows':[[c.text for c in r.cells] for r in t.rows]} for i,t in enumerate(d.tables,1)]
    if suffix in {'.txt','.md'}:
        return [{'block':1,'text':path.read_text(encoding='utf-8-sig')}]
    raise ValueError('Supported input: PDF, DOCX, TXT, MD')

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input',type=Path,required=True)
    parser.add_argument('--out',type=Path,required=True)
    args=parser.parse_args()
    try:
        if args.out.exists():raise FileExistsError('Refusing to overwrite existing lesson input')
        blocks=extract(args.input)
        if not blocks or not any(x.get('text','').strip() or x.get('rows') for x in blocks):
            raise ValueError('No readable text; use explicit OCR before lesson distillation')
        data={'source':str(args.input.resolve()),'extracted_at':datetime.now(timezone.utc).isoformat(),'status':'source_extracted_not_distilled','privacy':'source text may contain identities; author must approve any sharing','blocks':blocks,'lesson_draft':[]}
        args.out.parent.mkdir(parents=True,exist_ok=True)
        with args.out.open('x',encoding='utf-8') as f:json.dump(data,f,ensure_ascii=False,indent=2)
        print(json.dumps({'output':str(args.out.resolve()),'blocks':len(blocks),'status':data['status']},ensure_ascii=False))
        return 0
    except (ValueError,OSError) as exc:parser.exit(1,f'{exc}\n')
if __name__=='__main__':raise SystemExit(main())
