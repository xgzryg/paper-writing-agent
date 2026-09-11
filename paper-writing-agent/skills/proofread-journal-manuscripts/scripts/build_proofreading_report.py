#!/usr/bin/env python3
"""Build a landscape DOCX proofreading report from a validated JSON ledger."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


VALID_TYPES = {"Correction", "Query", "Author decision"}
VALID_CONFIDENCE = {"High", "Medium", "Low"}
VALID_SOURCE_FORMATS = {"PDF", "DOCX"}
VALID_SEGMENT_STATUS = {"verified", "uncertain"}
VALID_FINDING_BASIS = {"text", "visual"}
HEADERS = [
    "ID",
    "Type",
    "Current",
    "Proposed correction\nor query",
    "Reason",
    "Confidence",
]
COLUMN_WIDTHS = [0.50, 0.85, 2.15, 2.35, 3.25, 1.00]
CONTENT_WIDTH_DXA = round(sum(COLUMN_WIDTHS) * 1440)
INK = RGBColor(35, 37, 41)
MUTED = RGBColor(92, 98, 105)
RULE = "D9DDE3"
REPORT_FONT = None  # Optional explicitly selected installed font; otherwise use document/system fallback.


def set_font(run, size=10.5, bold=False, italic=False, color=INK):
    if REPORT_FONT:
        run.font.name = REPORT_FONT
        for slot in ("ascii", "hAnsi", "eastAsia"):
            run._element.get_or_add_rPr().rFonts.set(qn(f"w:{slot}"), REPORT_FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_style_font(style, size, *, bold=False, color=INK):
    if REPORT_FONT:
        style.font.name = REPORT_FONT
        r_pr = style._element.get_or_add_rPr()
        r_fonts = r_pr.rFonts
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.append(r_fonts)
        for slot in ("ascii", "hAnsi", "eastAsia"):
            r_fonts.set(qn(f"w:{slot}"), REPORT_FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color


def add_inline_markdown(paragraph, text, *, size=10.5, bold=False, color=INK):
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", str(text))
    for part in parts:
        if not part:
            continue
        part_bold = bold
        part_italic = False
        value = part
        if part.startswith("**") and part.endswith("**"):
            value = part[2:-2]
            part_bold = True
        elif part.startswith("*") and part.endswith("*"):
            value = part[1:-1]
            part_italic = True
        run = paragraph.add_run(value)
        set_font(run, size=size, bold=part_bold, italic=part_italic, color=color)


def visible_text(text):
    """Remove only the emphasis markers supported by add_inline_markdown."""
    value = str(text)
    value = re.sub(r"\*\*([^*\n]+)\*\*", r"\1", value)
    value = re.sub(r"\*([^*\n]+)\*", r"\1", value)
    return value


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is not None:
        tbl_pr.remove(borders)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "right", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    for edge in ("bottom", "insideH"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), RULE)
        borders.append(node)
    tbl_pr.append(borders)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in COLUMN_WIDTHS:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(round(width * 1440)))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell, width in zip(row.cells, COLUMN_WIDTHS):
            set_cell_width(cell, round(width * 1440))
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = paragraph.add_run("Page ")
    set_font(label, size=8.5, color=MUTED)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])
    set_font(run, size=8.5, color=MUTED)


def validate(data):
    required = {
        "source_file",
        "source_format",
        "article_type",
        "transcript_status",
        "coverage",
        "sections",
    }
    missing = required - set(data)
    if missing:
        raise ValueError(f"Missing top-level fields: {sorted(missing)}")
    if not isinstance(data["coverage"], list) or not isinstance(data["sections"], list):
        raise ValueError("coverage and sections must be lists")
    if data["source_format"] not in VALID_SOURCE_FORMATS:
        raise ValueError(f"Invalid source_format: {data['source_format']}")
    if data["transcript_status"] != "locked":
        raise ValueError("transcript_status must be 'locked' before report generation")

    if data["coverage"] != [section.get("title") for section in data["sections"]]:
        raise ValueError("coverage must exactly match section titles in source order")
    if not data["sections"]:
        raise ValueError("At least one included section is required")

    seen = set()
    seen_sections = set()
    seen_segments = set()
    numeric_ids = []
    for section in data["sections"]:
        for key in ("id", "title", "segments", "findings"):
            if key not in section:
                raise ValueError(f"Section missing {key}: {section}")
        if section["id"] in seen_sections:
            raise ValueError(f"Duplicate section ID: {section['id']}")
        seen_sections.add(section["id"])
        if not isinstance(section["segments"], list) or not section["segments"]:
            raise ValueError(f"segments must be a list in section {section['id']}")

        section_segments = {}
        for segment in section["segments"]:
            for key in ("id", "page_range", "text", "evidence_ref", "status"):
                if key not in segment:
                    raise ValueError(f"Transcript segment missing {key}: {segment}")
            segment_id = segment["id"]
            if segment_id in seen_segments:
                raise ValueError(f"Duplicate transcript segment ID: {segment_id}")
            seen_segments.add(segment_id)
            section_segments[segment_id] = segment
            if segment["status"] not in VALID_SEGMENT_STATUS:
                raise ValueError(
                    f"Invalid transcript status for {segment_id}: {segment['status']}"
                )
            if not isinstance(segment["text"], str) or not segment["text"]:
                raise ValueError(f"Transcript segment text must be non-empty: {segment_id}")
            if not isinstance(segment["evidence_ref"], str) or not segment["evidence_ref"].strip():
                raise ValueError(f"Transcript segment needs an evidence_ref: {segment_id}")
            if "display_text" in segment and visible_text(segment["display_text"]) != segment["text"]:
                raise ValueError(
                    f"display_text changes visible transcript characters: {segment_id}"
                )

        for finding in section["findings"]:
            for key in (
                "id",
                "type",
                "current",
                "proposed",
                "reason",
                "confidence",
                "segment_id",
                "source_excerpt",
                "basis",
            ):
                if key not in finding:
                    raise ValueError(f"Finding missing {key}: {finding}")
            fid = finding["id"]
            if not re.fullmatch(r"C\d{3,}", fid):
                raise ValueError(f"Invalid finding ID: {fid}")
            if fid in seen:
                raise ValueError(f"Duplicate finding ID: {fid}")
            seen.add(fid)
            numeric_ids.append(int(fid[1:]))
            if finding["type"] not in VALID_TYPES:
                raise ValueError(f"Invalid finding type: {finding['type']}")
            if finding["confidence"] not in VALID_CONFIDENCE:
                raise ValueError(f"Invalid confidence: {finding['confidence']}")
            segment_id = finding["segment_id"]
            if segment_id not in section_segments:
                raise ValueError(
                    f"Finding {fid} references a missing or cross-section segment: {segment_id}"
                )
            segment = section_segments[segment_id]
            if segment["status"] != "verified":
                raise ValueError(f"Finding {fid} is based on uncertain segment {segment_id}")
            source_excerpt = finding["source_excerpt"]
            if not isinstance(source_excerpt, str) or not source_excerpt:
                raise ValueError(f"Finding {fid} has an empty source_excerpt")
            if source_excerpt not in segment["text"]:
                raise ValueError(
                    f"Finding {fid} source_excerpt is not an exact substring of {segment_id}"
                )
            if visible_text(finding["current"]) != source_excerpt:
                raise ValueError(
                    f"Finding {fid} Current does not exactly represent source_excerpt"
                )
            if finding["basis"] not in VALID_FINDING_BASIS:
                raise ValueError(f"Invalid basis for {fid}: {finding['basis']}")
            if finding["basis"] == "visual":
                visual_evidence = finding.get("visual_evidence", "")
                if not isinstance(visual_evidence, str) or not visual_evidence.strip():
                    raise ValueError(f"Visual finding {fid} lacks visual_evidence")
    if numeric_ids and numeric_ids != list(range(1, len(numeric_ids) + 1)):
        raise ValueError("Finding IDs must be monotonically sequential from C001")


def configure_styles(doc):
    normal = doc.styles["Normal"]
    set_style_font(normal, 10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[style_name]
        set_style_font(style, size, bold=True)
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True


def add_metadata_line(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(f"{label}: ")
    set_font(run, size=9.5, bold=True, color=MUTED)
    add_inline_markdown(p, value, size=9.5, color=MUTED)


def build_report(data, output_path):
    validate(data)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)
    configure_styles(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("Proofreading Report")
    set_font(run, size=20, bold=True)

    counts = Counter(
        finding["type"]
        for section_data in data["sections"]
        for finding in section_data["findings"]
    )
    segment_counts = Counter(
        segment["status"]
        for section_data in data["sections"]
        for segment in section_data["segments"]
    )
    add_metadata_line(doc, "Source", data["source_file"])
    add_metadata_line(doc, "Source format", data["source_format"])
    add_metadata_line(doc, "Article type", data["article_type"])
    add_metadata_line(
        doc,
        "Transcript verification",
        f"{segment_counts['verified']} verified; {segment_counts['uncertain']} uncertain; locked",
    )
    add_metadata_line(doc, "Coverage", "; ".join(data.get("coverage", [])) or "None")
    add_metadata_line(doc, "Exclusions", "; ".join(data.get("exclusions", [])) or "None")
    add_metadata_line(doc, "Limitations", "; ".join(data.get("limitations", [])) or "None")
    add_metadata_line(
        doc,
        "Findings",
        f"{counts['Correction']} Correction; {counts['Query']} Query; "
        f"{counts['Author decision']} Author decision",
    )

    for section_data in data["sections"]:
        heading = doc.add_paragraph(style="Heading 1")
        add_inline_markdown(heading, section_data["title"], size=15, bold=True)
        findings = section_data["findings"]
        if not findings:
            p = doc.add_paragraph()
            add_inline_markdown(p, "No correction identified.", size=10.5, color=MUTED)
            continue

        table = doc.add_table(rows=1, cols=6)
        for cell, header in zip(table.rows[0].cells, HEADERS):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline_markdown(p, header, size=10, bold=True)
        set_repeat_table_header(table.rows[0])

        for finding in findings:
            row = table.add_row()
            values = [
                finding["id"],
                finding["type"],
                finding["current"],
                finding["proposed"],
                finding["reason"],
                finding["confidence"],
            ]
            for index, (cell, value) in enumerate(zip(row.cells, values)):
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                if index in (0, 5):
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline_markdown(p, value, size=10)
        set_table_geometry(table)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    doc.core_properties.title = "Proofreading Report"
    doc.core_properties.subject = f"Proofreading findings for {data['source_file']}"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "proofreading, academic manuscript, correction report"
    doc.save(output_path)


def main():
    global REPORT_FONT
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("ledger_json", type=Path)
    parser.add_argument("output_docx", type=Path)
    parser.add_argument("--font", help="Optional font actually installed on this target system")
    args = parser.parse_args()
    REPORT_FONT = args.font
    data = json.loads(args.ledger_json.read_text(encoding="utf-8"))
    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    build_report(data, args.output_docx)


if __name__ == "__main__":
    main()
