"""Small portable Word helpers; preserves original files and refuses output overwrite."""
import argparse
import copy
import datetime
import io
import json
from pathlib import Path
import shutil
import subprocess
import xml.etree.ElementTree as ET
import zipfile

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': W}


def fresh_write(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('xb') as f:
        f.write(data)


def extract(source, output):
    parts = []
    counts = dict.fromkeys(['drawing', 'object', 'fldChar', 'instrText', 'ins', 'del', 'hyperlink', 'oMath'], 0)
    with zipfile.ZipFile(source) as z:
        for name in z.namelist():
            if not (name == 'word/document.xml' or name.startswith(('word/header', 'word/footer', 'word/footnotes', 'word/endnotes', 'word/comments'))):
                continue
            if not name.endswith('.xml'):
                continue
            root = ET.fromstring(z.read(name))
            for node in root.iter():
                tag = node.tag.rsplit('}', 1)[-1]
                if tag in counts:
                    counts[tag] += 1
            paragraphs = []
            for index, para in enumerate(root.findall('.//w:p', NS), 1):
                text = ''.join(node.text or '' for node in para.iter() if node.tag in {f'{{{W}}}t', f'{{{W}}}delText'})
                paragraphs.append({'paragraph': index, 'text_including_revision_text': text})
            tables = []
            for table_index, table in enumerate(root.findall('.//w:tbl', NS), 1):
                rows = []
                for row in table.findall('w:tr', NS):
                    rows.append(['\n'.join(''.join(n.text or '' for n in p.findall('.//w:t', NS)) for p in cell.findall('w:p', NS)) for cell in row.findall('w:tc', NS)])
                tables.append({'table': table_index, 'rows': rows})
            parts.append({'part': name, 'paragraphs': paragraphs, 'tables': tables})
    result = {'source': str(source.resolve()), 'status': 'xml_content_only_not_layout_verified',
              'revision_note': 'Deleted and inserted text are both included; resolve intended revision view before quoting.',
              'object_counts': counts, 'parts': parts}
    fresh_write(output, json.dumps(result, ensure_ascii=False, indent=2).encode('utf-8'))


def from_json(source, output):
    from docx import Document
    from docx.shared import Pt
    data = json.loads(source.read_text(encoding='utf-8'))
    doc = Document()
    doc.core_properties.author = data.get('author', '')
    doc.core_properties.title = data.get('title', '')
    doc.styles['Normal'].font.size = Pt(11)
    if data.get('title'):
        doc.add_paragraph(data['title'], style='Title')
    for block in data['blocks']:
        kind = block.get('type', 'paragraph')
        if kind == 'heading':
            doc.add_heading(block['text'], level=int(block.get('level', 1)))
        elif kind == 'paragraph':
            doc.add_paragraph(block['text'])
        elif kind == 'table':
            rows = block['rows']
            if not rows or any(len(row) != len(rows[0]) for row in rows):
                raise ValueError('Table rows must have a consistent, nonzero shape')
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = 'Table Grid'
            for row, values in zip(table.rows, rows):
                for cell, value in zip(row.cells, values):
                    cell.text = str(value)
        else:
            raise ValueError(f'Unsupported block type {kind}')
    buffer = io.BytesIO()
    doc.save(buffer)
    fresh_write(output, buffer.getvalue())


def replace(source, edits_file, output, track, author):
    # lxml preserves namespace prefixes used in Word compatibility attributes.
    from lxml import etree
    data = json.loads(edits_file.read_text(encoding='utf-8'))
    with zipfile.ZipFile(source) as z:
        root = etree.fromstring(z.read('word/document.xml'))
        ids = [int(x) for x in root.xpath('//@w:id', namespaces=NS) if x.isdigit()]
        change_id = max(ids, default=0) + 1
        for change in data['replacements']:
            old, new = change['old'], change['new']
            if not old:
                raise ValueError('Empty old text is unsupported')
            candidates = [n for n in root.findall('.//w:t', NS) if old in (n.text or '')]
            if sum((n.text or '').count(old) for n in candidates) != 1:
                raise ValueError(f'Replacement must uniquely match one text run: {old!r}')
            node = candidates[0]
            if not track:
                node.text = node.text.replace(old, new, 1)
                node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                continue
            run = node.getparent()
            if run.tag != f'{{{W}}}r' or any(a.tag in {f'{{{W}}}ins', f'{{{W}}}del'} for a in run.iterancestors()):
                raise ValueError('Tracked replacement inside an existing revision is unsupported')
            non_props = [c for c in run if c.tag != f'{{{W}}}rPr']
            if non_props != [node]:
                raise ValueError('Tracked replacement requires a simple text run; preserve complex objects manually')
            before, after = node.text.split(old, 1)
            parent, index = run.getparent(), run.getparent().index(run)
            style = run.find('w:rPr', NS)
            def new_run(value, deleted=False):
                r = etree.Element(f'{{{W}}}r')
                if style is not None:
                    r.append(copy.deepcopy(style))
                t = etree.SubElement(r, f'{{{W}}}' + ('delText' if deleted else 't'))
                t.text = value
                t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                return r
            replacements = []
            if before:
                replacements.append(new_run(before))
            for kind, value in [('del', old), ('ins', new)]:
                if value:
                    element = etree.Element(f'{{{W}}}{kind}')
                    element.set(f'{{{W}}}id', str(change_id))
                    element.set(f'{{{W}}}author', author)
                    element.set(f'{{{W}}}date', datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'))
                    change_id += 1
                    element.append(new_run(value, kind == 'del'))
                    replacements.append(element)
            if after:
                replacements.append(new_run(after))
            parent.remove(run)
            for offset, element in enumerate(replacements):
                parent.insert(index + offset, element)
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w', compression=zipfile.ZIP_DEFLATED) as out:
            for info in z.infolist():
                contents = etree.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True) if info.filename == 'word/document.xml' else z.read(info)
                out.writestr(info, contents)
    fresh_write(output, buffer.getvalue())


def render(source, output_dir, renderer):
    executable = renderer or shutil.which('soffice') or shutil.which('libreoffice')
    if not executable:
        raise RuntimeError('No LibreOffice renderer found. Supply --renderer or use an available Word export tool; layout is unverified.')
    if (output_dir / (source.stem + '.pdf')).exists():
        raise FileExistsError('PDF output already exists; choose a new task render directory')
    output_dir.mkdir(parents=True, exist_ok=True)
    profile = (output_dir / 'libreoffice-profile').resolve()
    result = subprocess.run([str(executable), f'-env:UserInstallation={profile.as_uri()}', '--headless', '--convert-to', 'pdf', '--outdir', str(output_dir.resolve()), str(source.resolve())], capture_output=True, text=True, timeout=120)
    print(result.stdout)
    if result.returncode or not (output_dir / (source.stem + '.pdf')).exists():
        raise RuntimeError(f'Render failed; layout remains unverified: {result.stderr}')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    for name in ['extract', 'from-json', 'replace']:
        p = sub.add_parser(name)
        p.add_argument('input', type=Path)
        p.add_argument('--output', type=Path, required=True)
        if name == 'replace':
            p.add_argument('--edits', type=Path, required=True)
            p.add_argument('--track', action='store_true')
            p.add_argument('--author', default='Editor')
    p = sub.add_parser('render')
    p.add_argument('input', type=Path)
    p.add_argument('--output-dir', type=Path, required=True)
    p.add_argument('--renderer')
    args = parser.parse_args()
    if args.command == 'extract':
        extract(args.input, args.output)
    elif args.command == 'from-json':
        from_json(args.input, args.output)
    elif args.command == 'replace':
        replace(args.input, args.edits, args.output, args.track, args.author)
    else:
        render(args.input, args.output_dir, args.renderer)
    print('Completed file operation; visual verification is a separate step.')


if __name__ == '__main__':
    main()
