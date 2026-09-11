#!/usr/bin/env python3
"""Render a grounded review/response payload to synchronized DOCX and Markdown.

This renderer formats supplied content; it does not evaluate scientific truth,
retrieve historical PRF cases, or launch reviewers.
"""
from __future__ import annotations
import argparse
from io import BytesIO
import json
from pathlib import Path

STATUSES={'DONE','VERIFIED_DONE','DRAFTED','REPORTED_DONE_UNVERIFIED','TODO_TEXT','TODO_ANALYSIS','TODO_EXPERIMENT','TODO_AUTHOR_CONFIRM','NOT_FEASIBLE','PROPOSED_DISAGREEMENT'}

def require_text(value, field):
    if not isinstance(value,str) or not value.strip():
        raise ValueError(f'{field} must be a nonempty string')
    return value.strip()

def source_text(path):
    if path.suffix.lower()=='.docx':
        from docx import Document
        doc=Document(path)
        return '\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
    if path.suffix.lower()=='.pdf':
        from pypdf import PdfReader
        return '\n'.join(p.extract_text() or '' for p in PdfReader(path).pages)
    return path.read_text(encoding='utf-8-sig')

def validate(payload, payload_dir):
    if payload.get('entry') not in {'review','respond'}:
        raise ValueError('entry must be review or respond')
    if payload.get('language') not in {'zh','en'}:
        raise ValueError('language must be zh or en')
    require_text(payload.get('case_id'),'case_id')
    require_text(payload.get('scope'),'scope')
    reviewers=payload.get('reviewers')
    if not isinstance(reviewers,list) or not reviewers:
        raise ValueError('reviewers must be a nonempty list')
    seen=set()
    for ri, reviewer in enumerate(reviewers,1):
        require_text(reviewer.get('label'),f'reviewers[{ri}].label')
        require_text(reviewer.get('overall'),f'reviewers[{ri}].overall')
        for tier in ['major','minor']:
            concerns=reviewer.get(tier)
            if not isinstance(concerns,list):
                raise ValueError(f'reviewers[{ri}].{tier} must be a list; [] is allowed')
            for item in concerns:
                for field in ['id','claim','evidence','body','resolution']:
                    require_text(item.get(field),field)
                if item['id'] in seen:
                    raise ValueError('duplicate concern ID: '+item['id'])
                seen.add(item['id'])
                if not isinstance(item.get('blocking'),bool):
                    raise ValueError('blocking must be boolean')
                if tier=='minor' and item['blocking']:
                    raise ValueError('Minor concerns cannot be blocking')
                if payload['entry']=='respond':
                    for field in ['response','changes','location','author_action']:
                        require_text(item.get(field),field)
                    if item.get('status') not in STATUSES:
                        raise ValueError('invalid response status')
                    if item['status'] in {'DONE','VERIFIED_DONE'}:
                        evidence=item.get('verification',{})
                        source=Path(require_text(evidence.get('path'),'verification.path'))
                        if not source.is_absolute():
                            source=payload_dir/source
                        excerpt=require_text(evidence.get('excerpt'),'verification.excerpt')
                        if not source.is_file() or excerpt not in source_text(source):
                            raise ValueError('verified completion excerpt not found in the supplied artifact')
    if not isinstance(payload.get('synthesis',[]),list):
        raise ValueError('synthesis must be a list')
    for item in payload.get('synthesis',[]):
        require_text(item.get('text'),'synthesis.text')
        ids=item.get('concern_ids',[])
        if not ids or any(i not in seen for i in ids):
            raise ValueError('synthesis must reference existing concern IDs')
    rows=payload.get('tasks',[])
    if not isinstance(rows,list):
        raise ValueError('tasks must be a list')
    for row in rows:
        if not isinstance(row,list) or len(row)!=8 or any(not isinstance(x,str) for x in row):
            raise ValueError('each task row must contain eight strings')

def blocks(payload):
    zh=payload['language']=='zh'
    b=[('h1', '论文审稿报告' if zh and payload['entry']=='review' else '返修工作报告' if zh else 'Manuscript review' if payload['entry']=='review' else 'Revision work report'),
       ('p',payload['case_id']),('h2','审稿依据' if zh else 'Material scope'),('p',payload['scope'])]
    if payload.get('assessment'):
        b += [('h2','模拟评估' if zh else 'Simulated assessment'),('p',str(payload['assessment']))]
    fields=[('claim','论断' if zh else 'Claim'),('evidence','证据位置' if zh else 'Evidence'),('body','意见' if zh else 'Concern'),('resolution','解决标准' if zh else 'Resolution test')]
    for reviewer in payload['reviewers']:
        b += [('h2',reviewer['label']),('p',reviewer['overall'])]
        if reviewer.get('strengths'):
            b += [('p',('优点：' if zh else 'Strengths: ')+str(reviewer['strengths']))]
        for tier,title in [('major','主要问题' if zh else 'Major concerns'),('minor','次要问题' if zh else 'Minor comments')]:
            b.append(('h3',title))
            if not reviewer[tier]:
                b.append(('p','依据可见材料未发现该级别问题。' if zh else 'None identified from the supplied material.'))
            for item in reviewer[tier]:
                b += [('h4',item['id']),('p',('阻塞：' if zh else 'Blocking: ')+('Yes' if item['blocking'] else 'No'))]
                b += [('p',label+'：'+item[field] if zh else label+': '+item[field]) for field,label in fields]
                if payload['entry']=='respond':
                    b += [('p', ('状态：' if zh else 'Status: ')+item['status'])]
                    for field,label in [('response','回应' if zh else 'Response'),('changes','实际修改' if zh else 'Concrete revisions'),('location','修订位置' if zh else 'Location'),('author_action','作者待办' if zh else 'Author action')]:
                        b.append(('p',label+'：'+item[field] if zh else label+': '+item[field]))
    if payload.get('synthesis'):
        b.append(('h2','作者/编辑侧汇总（不传给审稿人）' if zh else 'Author/editor synthesis (not reviewer-facing)'))
        for item in payload['synthesis']:
            b.append(('p',item['text']+' ['+', '.join(item['concern_ids'])+']'))
    if payload.get('tasks'):
        b.append(('h2','修订任务' if zh else 'Revision tasks'))
        b.append(('table',[['ID','Reviewer','Concern','Strategy','Status','Input needed','Output','Blocks response?']]+payload['tasks']))
    return b

def render(payload, out, payload_dir):
    validate(payload,payload_dir)
    if out.suffix.lower()!='.docx':
        raise ValueError('--out must end in .docx')
    md=out.with_suffix('.md')
    if out.exists() or md.exists():
        raise FileExistsError('Refusing to overwrite an existing DOCX or Markdown file')
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    doc=Document()
    section=doc.sections[0]
    section.page_height=Cm(29.7); section.page_width=Cm(21)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Cm(2.5)
    doc.styles['Normal'].font.size=Pt(10.5)
    doc.styles['Normal'].font.color.rgb=RGBColor(0,0,0)
    doc.styles['Normal'].paragraph_format.line_spacing=1.15
    for level,size in [(1,16),(2,13),(3,11),(4,10.5)]:
        style=doc.styles['Heading '+str(level)]
        style.font.size=Pt(size)
        style.font.color.rgb=RGBColor(0,0,0)
    markdown=[]
    for kind,value in blocks(payload):
        if kind.startswith('h'):
            level=int(kind[1:]);doc.add_heading(value,level=level)
            markdown.extend(['#'*level+' '+value,''])
        elif kind=='p':
            doc.add_paragraph(value);markdown.extend([value,''])
        else:
            table=doc.add_table(rows=0,cols=len(value[0]));table.style='Table Grid'
            for row in value:
                cells=table.add_row().cells
                for c,t in zip(cells,row):c.text=t
            for i,row in enumerate(value):
                markdown.append('| '+' | '.join(t.replace('|','\\|').replace('\n','<br>') for t in row)+' |')
                if i==0:markdown.append('| '+' | '.join(['---']*len(row))+' |')
            markdown.append('')
    memory=BytesIO();doc.save(memory)
    out.parent.mkdir(parents=True,exist_ok=True)
    with out.open('xb') as f:f.write(memory.getvalue())
    with md.open('x',encoding='utf-8') as f:f.write('\n'.join(markdown))
    return {'docx':str(out.resolve()),'markdown':str(md.resolve()),'visual_qa':'not_performed','verification':'payload structure and cited completion excerpts checked; scientific assessment requires reviewer judgement'}

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--payload',type=Path,required=True)
    p.add_argument('--out',type=Path,required=True)
    p.add_argument('--language',choices=['zh','en'])
    args=p.parse_args()
    try:
        payload=json.loads(args.payload.read_text(encoding='utf-8-sig'))
        if args.language:payload['language']=args.language
        print(json.dumps(render(payload,args.out,args.payload.resolve().parent),ensure_ascii=False,indent=2))
        return 0
    except (ValueError,OSError,KeyError,TypeError) as exc:
        p.exit(1,f'{exc}\n')

if __name__=='__main__':
    raise SystemExit(main())
